"""
Limpieza y reporte de inventario.
Uso:  python limpiar_inventario.py inventario_sucio.xlsx
Genera: reporte_inventario.xlsx
"""
import sys
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ARCHIVO = sys.argv[1] if len(sys.argv) > 1 else "inventario_sucio.xlsx"
log = []  # registro de todo lo que se corrigió

df = pd.read_excel(ARCHIVO)
filas_inicio = len(df)

# 1) Duplicados exactos --------------------------------------------------
n = df.duplicated().sum()
df = df.drop_duplicates().copy()
log.append(("Filas duplicadas eliminadas", int(n)))

# 2) Texto: quitar espacios y unificar mayúsculas ------------------------
for col in ["Producto", "Categoria", "Proveedor", "Ubicacion"]:
    df[col] = df[col].astype("string").str.strip()
antes = df["Categoria"].copy()
# "EPP" es sigla: se respeta; el resto queda como "Primera mayúscula"
SIGLAS = {"epp": "EPP"}
df["Categoria"] = df["Categoria"].str.lower().map(lambda c: SIGLAS.get(c, c.capitalize()) if pd.notna(c) else c)
log.append(("Categorías con formato corregido", int((antes != df["Categoria"]).sum())))

# 3) Proveedor vacío o "N/A" ---------------------------------------------
vacios = df["Proveedor"].isna() | df["Proveedor"].isin(["", "N/A"])
df.loc[vacios, "Proveedor"] = "Sin proveedor"
log.append(("Proveedores vacíos marcados 'Sin proveedor'", int(vacios.sum())))

# 4) Fechas en formatos mezclados ----------------------------------------
def a_fecha(valor):
    if pd.isna(valor):
        return pd.NaT
    if isinstance(valor, pd.Timestamp):
        return valor
    texto = str(valor).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%y", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return pd.to_datetime(texto, format=fmt)
        except ValueError:
            continue
    return pd.NaT

fechas = df["Fecha_Ingreso"].apply(a_fecha)
log.append(("Fechas no reconocidas (revisar a mano)", int(fechas.isna().sum())))
df["Fecha_Ingreso"] = fechas

# 5) Costo: "$1,850" -> 1850 ----------------------------------------------
costo = (df["Costo_Unitario"].astype(str)
         .str.replace(r"[$,\s]", "", regex=True))
df["Costo_Unitario"] = pd.to_numeric(costo, errors="coerce")
log.append(("Costos sin valor válido", int(df["Costo_Unitario"].isna().sum())))

# 6) Stock: texto, vacíos y negativos ------------------------------------
stock_original = df["Stock_Actual"].copy()
solo_numeros = stock_original.astype(str).str.extract(r"(-?\d+)")[0]
df["Stock_Actual"] = pd.to_numeric(solo_numeros, errors="coerce")
log.append(("Stock vacío o ilegible (revisar)", int(df["Stock_Actual"].isna().sum())))
neg = df["Stock_Actual"] < 0
log.append(("Stock negativo (error de carga, revisar)", int(neg.sum())))
df["Revisar"] = ""
df.loc[df["Stock_Actual"].isna() | neg, "Revisar"] = "Stock inválido"
df.loc[df["Costo_Unitario"].isna(), "Revisar"] += " Costo inválido"
df.loc[df["Fecha_Ingreso"].isna(), "Revisar"] += " Fecha inválida"
df["Revisar"] = df["Revisar"].str.strip()

# 7) Cálculos del reporte -------------------------------------------------
limpio = df[df["Revisar"] == ""].copy()
limpio["Stock_Actual"] = limpio["Stock_Actual"].astype(int)
limpio["Valor_Stock"] = limpio["Stock_Actual"] * limpio["Costo_Unitario"]
limpio["Estado"] = limpio.apply(
    lambda r: "SIN STOCK" if r.Stock_Actual == 0
    else ("STOCK BAJO" if r.Stock_Actual < r.Stock_Minimo else "OK"), axis=1)

alertas = limpio[limpio["Estado"] != "OK"].sort_values(["Estado", "Stock_Actual"], ascending=[False, True])
resumen = (limpio.groupby("Categoria")
           .agg(Productos=("Codigo", "count"), Unidades=("Stock_Actual", "sum"),
                Valor_Total=("Valor_Stock", "sum"),
                En_Alerta=("Estado", lambda s: (s != "OK").sum()))
           .reset_index().sort_values("Valor_Total", ascending=False))
revisar = df[df["Revisar"] != ""]

log += [("Filas al inicio", filas_inicio), ("Filas limpias en el reporte", len(limpio)),
        ("Filas enviadas a revisión manual", len(revisar))]
log_df = pd.DataFrame(log, columns=["Control", "Cantidad"])

# 8) Escribir Excel con formato ------------------------------------------
SALIDA = "reporte_inventario.xlsx"
hojas = {"Resumen": resumen, "Alertas": alertas, "Inventario_Limpio": limpio,
         "A_Revisar": revisar, "Log_Limpieza": log_df}

with pd.ExcelWriter(SALIDA, engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
    for nombre, datos in hojas.items():
        datos.to_excel(w, sheet_name=nombre, index=False)
        ws = w.sheets[nombre]
        for celda in ws[1]:
            celda.font = Font(name="Arial", bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
            celda.alignment = Alignment(horizontal="center")
        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.font = Font(name="Arial")
        for i, col in enumerate(datos.columns, 1):
            ancho = max(len(str(col)), *(len(str(v)) for v in datos[col])) if len(datos) else len(str(col))
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 3, 40)
            if col in ("Costo_Unitario", "Valor_Stock", "Valor_Total"):
                for c in ws[get_column_letter(i)][1:]:
                    c.number_format = '$#,##0'
        ws.freeze_panes = "A2"
        if nombre in ("Alertas", "Inventario_Limpio"):
            ws.auto_filter.ref = ws.dimensions
        if nombre == "Alertas":
            rojo = PatternFill("solid", fgColor="F8CBAD")
            amarillo = PatternFill("solid", fgColor="FFF2CC")
            col_estado = list(datos.columns).index("Estado") + 1
            for fila in ws.iter_rows(min_row=2):
                fila[col_estado - 1].fill = rojo if fila[col_estado - 1].value == "SIN STOCK" else amarillo

print(log_df.to_string(index=False))
print(f"\nReporte generado: {SALIDA}")
